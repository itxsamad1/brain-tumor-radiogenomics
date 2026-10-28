#!/usr/bin/env python3
"""Build a paper-style Word manuscript for Paper A from experiment artifacts."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EXP = ROOT / "experiments" / "paper_a"
OUT = EXP / "Paper_A_Leakage_Shortcuts_Manuscript.docx"


def set_run_font(run, size=11, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(doc, text, *, size=11, bold=False, italic=False, space_after=8, align="justify", first_line=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line and align == "justify":
        pf.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=10, italic=True)


def add_figure(doc, path: Path, caption: str, width=5.8):
    if not path.exists():
        add_para(doc, f"[Figure missing: {path.name}]", italic=True, first_line=False, align="center")
        add_caption(doc, caption)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_table(doc, headers, rows, caption: str):
    add_caption(doc, caption)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def main() -> int:
    summary = load_json(EXP / "summary.json")
    fold_metrics = {f: load_json(EXP / f"fold{f}_resnet18" / "metrics.json") for f in range(1, 6)}
    leaky = load_json(EXP / "leaky_slice_split" / "metrics.json")
    leak_stats = load_json(EXP / "leaky_slice_split" / "leakage_stats.json")
    masked_eval = load_json(EXP / "tumor_masked" / "eval_masked_metrics.json")
    masked_train = load_json(EXP / "tumor_masked" / "train_masked_metrics.json")

    pl = summary["patient_level_resnet18"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(6)
    run = t.add_run(
        "Reality Check: Data Leakage and Imaging Shortcuts in "
        "2D Brain-Tumor MRI Classification on Figshare-Style Benchmarks"
    )
    set_run_font(run, size=16, bold=True)

    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = authors.add_run("brain-tumor-radiogenomics — Paper A experimental manuscript draft")
    set_run_font(r, size=11, italic=True)

    aff = doc.add_paragraph()
    aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aff.paragraph_format.space_after = Pt(18)
    r = aff.add_run(
        "Working draft generated from frozen experiment artifacts "
        f"({OUT.name}). Not a camera-ready journal submission."
    )
    set_run_font(r, size=9, italic=True)
    r.font.color.rgb = RGBColor(80, 80, 80)

    # Abstract
    add_heading_custom(doc, "Abstract", 1)
    add_para(
        doc,
        "High accuracies (often 98–99%) are routinely reported for convolutional neural networks "
        "on the Jun Cheng Figshare 2D brain-tumor MRI dataset. We ask how much of that performance "
        "reflects genuine patient-level generalization versus data leakage and non-tumor shortcuts. "
        "Using the same ResNet-18 recipe, we compare official patient-level five-fold cross-validation "
        "against an intentionally leaky slice-level random split, audit near-duplicate slices with "
        "perceptual hashing, evaluate Grad-CAM explanations, and run tumor-masked controls that zero "
        "the annotated tumor region. Patient-level ResNet-18 achieves mean macro-F1 "
        f"{pl['macro_f1_mean']:.3f} ± {pl['macro_f1_std']:.3f} and accuracy "
        f"{pl['accuracy_mean']:.3f} ± {pl['accuracy_std']:.3f}. The leaky slice-level protocol reaches "
        f"macro-F1 {summary['slice_level_leaky_macro_f1']:.3f} while "
        f"{100 * leak_stats['frac_val_slices_from_shared_patients']:.1f}% of validation slices share "
        "patients with the training set. Zeroing tumors at evaluation drops fold-1 macro-F1 from 0.962 "
        f"to {masked_eval['macro_f1']:.3f}, yet training from scratch on tumor-zeroed images still "
        f"recovers macro-F1 {masked_train['best_macro_f1']:.3f}. We conclude that patient-level splits "
        "are necessary, that slice-level reporting systematically inflates scores, and that Figshare-style "
        "2D benchmarks remain vulnerable to background shortcuts even when models attend to tumor tissue.",
        first_line=False,
    )
    add_para(
        doc,
        "Keywords: brain tumor MRI; data leakage; patient-level cross-validation; shortcut learning; "
        "Grad-CAM; Figshare; trustworthy medical AI.",
        italic=True,
        first_line=False,
        space_after=14,
    )

    # 1 Intro
    add_heading_custom(doc, "1. Introduction", 1)
    add_para(
        doc,
        "Public 2D MRI collections such as the Jun Cheng Figshare brain-tumor dataset "
        "(3,064 contrast-enhanced T1-weighted slices from 233 patients; meningioma, glioma, pituitary) "
        "have become a default playground for deep classification papers. Many report near-perfect "
        "accuracy. Reviewers increasingly ask whether those numbers arise from leakage—especially "
        "slice-level splits that place different slices of the same patient in both train and test—"
        "or from shortcuts unrelated to tumor morphology (imaging plane, skull outline, intensity "
        "context outside the lesion).",
    )
    add_para(
        doc,
        "This manuscript packages Paper A of the brain-tumor-radiogenomics program: a methodology / "
        "meta-science study that freezes patient-level folds derived from the official cvind indices, "
        "trains a transparent ResNet-18 baseline, and stress-tests the claim surface with audits. "
        "We deliberately do not position a new state-of-the-art classifier as the contribution. "
        "The contribution is a quantified reality check and a reproducible experiment package.",
    )
    add_para(
        doc,
        "Our research question is: how much of the widely reported 98–99% accuracy on Figshare-style "
        "2D brain-tumor classification comes from data leakage and shortcuts versus genuine "
        "patient-level generalization?",
    )

    # 2 Related
    add_heading_custom(doc, "2. Related work and split protocols", 1)
    add_para(
        doc,
        "Prior Figshare studies span classical CNNs, CapsNets, hybrid CNN–SVM pipelines, and "
        "attention variants. Reporting practices vary widely: some cite patient-wise folds from cvind, "
        "while many describe random slice hold-outs without documenting patient overlap. Table 4 "
        "summarizes a compact literature snapshot used for discussion; it is not an exhaustive "
        "systematic review.",
    )
    add_table(
        doc,
        ["Study (illustrative)", "Typical claim", "Split practice (as commonly reported)"],
        [
            ["Cheng et al. (dataset)", "Baseline resource", "Official cvind patient folds"],
            ["Many CNN / CapsNet papers", "≈98–99% accuracy", "Often slice-level or unclear"],
            ["This work (Paper A)", "Honest ~95% macro-F1", "Frozen patient-level 5-fold + leaky contrast"],
        ],
        "Table 4. Illustrative comparison of Figshare split practices (literature snapshot).",
    )

    # 3 Methods
    add_heading_custom(doc, "3. Methods", 1)
    add_heading_custom(doc, "3.1 Dataset", 2)
    add_para(
        doc,
        "We use the Jun Cheng Figshare collection (DOI 10.6084/m9.figshare.1512427; CC BY 4.0): "
        "3,064 T1c slices, 233 patients, three classes (1=meningioma, 2=glioma, 3=pituitary), "
        "with per-slice tumor masks. A packaged PNG export and frozen splits CSV were used for "
        "GPU training on Kaggle (`itxsamad/btlab-figshare-paper-a`).",
    )
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["Slices", "3,064"],
            ["Patients", "233"],
            ["Classes", "meningioma / glioma / pituitary"],
            ["Modality", "2D T1-weighted contrast-enhanced MRI"],
            ["Primary split", "Patient-level 5-fold (cvind-majority)"],
            ["Comparison split", "Slice-level stratified 80/20 (leaky)"],
        ],
        "Table 1. Dataset summary.",
    )

    add_heading_custom(doc, "3.2 Splits", 2)
    add_para(
        doc,
        "Primary protocol: patient-level five-fold cross-validation. No patient appears in more than "
        "one fold’s validation set. Comparison protocol: intentionally leaky slice-level random split "
        f"(seed 42, 20% validation, stratified by label). Under the leaky protocol, "
        f"{leak_stats['n_patients_in_both_splits']} patients appear in both train and validation, and "
        f"{100 * leak_stats['frac_val_slices_from_shared_patients']:.1f}% of validation slices come from "
        "those shared patients.",
    )

    add_heading_custom(doc, "3.3 Model and training", 2)
    add_para(
        doc,
        "Architecture: ResNet-18 with ImageNet initialization; final fully connected layer remapped "
        "to three classes. Images are resized to 224×224, grayscale expanded to three channels, "
        "normalized to mean/std 0.5. Training uses AdamW (lr 1e-4, weight decay 1e-4), cross-entropy, "
        "random horizontal flip and ±15° rotation on the training set, up to 30 epochs on Kaggle GPU "
        "with resumable checkpoint_last / checkpoint_best. Primary seed: 42. Fold isolation was "
        "enforced so multi-fold jobs cannot resume from another fold’s checkpoint.",
    )

    add_heading_custom(doc, "3.4 Metrics", 2)
    add_para(
        doc,
        "We report accuracy, macro-F1 (primary), balanced accuracy, per-class precision/recall/F1, "
        "one-vs-rest AUROC and AUPRC, expected calibration error (ECE) with reliability diagrams, "
        "and bootstrap 95% confidence intervals on macro-F1 and accuracy for each fold evaluation.",
    )

    add_heading_custom(doc, "3.5 Audits", 2)
    add_para(
        doc,
        "(1) Split audit — same model on patient-level vs leaky slice-level protocols. "
        "(2) Near-duplicate audit — difference hash (dHash) across patients. "
        "(3) Tumor-masked control — evaluate a patient-level checkpoint with tumor pixels zeroed; "
        "additionally train from scratch on tumor-zeroed images for 10 epochs. "
        "(4) Grad-CAM on correct and incorrect validation predictions (ResNet layer4).",
    )

    # 4 Results
    add_heading_custom(doc, "4. Results", 1)
    add_heading_custom(doc, "4.1 Patient-level five-fold performance", 2)
    add_para(
        doc,
        f"Across five patient-level folds, ResNet-18 obtains mean macro-F1 "
        f"{pl['macro_f1_mean']:.3f} ± {pl['macro_f1_std']:.3f} and mean accuracy "
        f"{pl['accuracy_mean']:.3f} ± {pl['accuracy_std']:.3f}. Fold-wise metrics are listed in Table 2. "
        "AUROC remains high (≈0.989–0.998) while ECE stays modest (≈0.017–0.043), indicating strong "
        "ranking performance with reasonably calibrated probabilities on this clean 2D task.",
        first_line=False,
    )

    fold_rows = []
    for f in range(1, 6):
        m = fold_metrics[f]
        fold_rows.append([
            str(f),
            f"{m['accuracy']:.3f}",
            f"{m['macro_f1']:.3f}",
            f"{m['auroc_macro_ovr']:.3f}",
            f"{m['ece']:.3f}",
            f"{m['bootstrap_macro_f1']['low']:.3f}–{m['bootstrap_macro_f1']['high']:.3f}",
        ])
    fold_rows.append([
        "Mean±SD",
        f"{pl['accuracy_mean']:.3f}±{pl['accuracy_std']:.3f}",
        f"{pl['macro_f1_mean']:.3f}±{pl['macro_f1_std']:.3f}",
        "—",
        "—",
        "—",
    ])
    add_table(
        doc,
        ["Fold", "Accuracy", "Macro-F1", "AUROC", "ECE", "Macro-F1 95% CI"],
        fold_rows,
        "Table 2. Patient-level ResNet-18 five-fold results.",
    )

    m1 = fold_metrics[1]
    add_para(
        doc,
        f"Fold 1 per-class F1 scores: meningioma {m1['per_class']['meningioma']['f1']:.3f}, "
        f"glioma {m1['per_class']['glioma']['f1']:.3f}, "
        f"pituitary {m1['per_class']['pituitary']['f1']:.3f}. "
        "Representative fold-1 figures follow.",
        first_line=False,
    )

    add_figure(
        doc,
        EXP / "fold1_resnet18" / "figures" / "confusion_matrix.png",
        "Figure 4. Fold-1 confusion matrix (patient-level ResNet-18).",
        width=4.5,
    )
    add_figure(
        doc,
        EXP / "fold1_resnet18" / "figures" / "roc.png",
        "Figure 5a. Fold-1 one-vs-rest ROC curves.",
        width=5.0,
    )
    add_figure(
        doc,
        EXP / "fold1_resnet18" / "figures" / "pr.png",
        "Figure 5b. Fold-1 precision–recall curves.",
        width=5.0,
    )
    add_figure(
        doc,
        EXP / "fold1_resnet18" / "figures" / "reliability.png",
        "Figure 6. Fold-1 reliability / calibration diagram.",
        width=5.0,
    )
    add_figure(
        doc,
        EXP / "fold1_resnet18" / "figures" / "training_curves.png",
        "Figure 7. Fold-1 training curves (loss and macro-F1).",
        width=5.2,
    )

    add_heading_custom(doc, "4.2 Leaky slice-level protocol", 2)
    add_para(
        doc,
        f"Under slice-level random splitting, the same architecture reaches accuracy "
        f"{leaky['accuracy']:.3f} and macro-F1 {leaky['macro_f1']:.3f} "
        f"(AUROC {leaky['auroc_macro_ovr']:.3f}, ECE {leaky['ece']:.3f}). "
        "The absolute macro-F1 gap versus the patient-level mean is approximately "
        f"{summary['slice_level_leaky_macro_f1'] - pl['macro_f1_mean']:.3f}. "
        "Because nearly all validation slices share patients with training, this gap is best read as "
        "leakage-driven inflation rather than an architectural improvement.",
    )
    add_figure(
        doc,
        EXP / "figures" / "patient_vs_slice_macro_f1.png",
        "Figure 3. Macro-F1 comparison: patient-level mean vs leaky slice-level vs tumor-zeroed evaluation.",
        width=5.2,
    )

    add_heading_custom(doc, "4.3 Near-duplicate audit", 2)
    add_para(
        doc,
        "Difference hashing over all 3,064 slices found 3,050 unique hashes and five cross-patient "
        "identical-hash groups. Near-duplicates therefore exist but are limited; they do not alone "
        "explain the larger inflation from slice-level splitting, which is dominated by intra-patient "
        "slice reuse across splits.",
    )

    add_heading_custom(doc, "4.4 Tumor-masked control", 2)
    add_para(
        doc,
        f"Evaluating the fold-1 patient-level checkpoint on images with tumor pixels zeroed yields "
        f"accuracy {masked_eval['accuracy']:.3f} and macro-F1 {masked_eval['macro_f1']:.3f} "
        f"(drop of {0.9615053504446985 - masked_eval['macro_f1']:.3f} macro-F1 relative to unmasked "
        "fold-1). Meningioma recall collapses most severely, consistent with reliance on lesion "
        f"appearance. However, training from scratch for 10 epochs on tumor-zeroed images still "
        f"reaches best macro-F1 {masked_train['best_macro_f1']:.3f}—nearly matching the honest "
        "patient-level mean. Models therefore both use tumor evidence and exploit strong non-tumor "
        "context when forced to.",
    )
    add_table(
        doc,
        ["Protocol", "Macro-F1", "Accuracy", "Interpretation"],
        [
            [
                "Patient-level (5-fold mean)",
                f"{pl['macro_f1_mean']:.3f}",
                f"{pl['accuracy_mean']:.3f}",
                "Primary honest estimate",
            ],
            [
                "Slice-level leaky",
                f"{summary['slice_level_leaky_macro_f1']:.3f}",
                f"{summary['slice_level_leaky_accuracy']:.3f}",
                "Inflated by patient overlap",
            ],
            [
                "Tumor-zeroed eval (fold 1)",
                f"{masked_eval['macro_f1']:.3f}",
                f"{masked_eval['accuracy']:.3f}",
                "Uses tumor region",
            ],
            [
                "Tumor-zeroed train (fold 1)",
                f"{masked_train['best_macro_f1']:.3f}",
                "—",
                "Background shortcuts remain",
            ],
        ],
        "Table 3. Ablation / audit summary.",
    )

    add_heading_custom(doc, "4.5 Grad-CAM examples", 2)
    add_para(
        doc,
        "Grad-CAM overlays on fold-1 validation slices show attention concentrated near lesions for "
        "many correct predictions, while some failures highlight peripheral structure—qualitatively "
        "aligned with the masked-control evidence for mixed tumor and context use.",
        first_line=False,
    )
    for name, cap in [
        ("correct_0.png", "Figure 8a. Grad-CAM — correct prediction example."),
        ("wrong_0.png", "Figure 8b. Grad-CAM — incorrect prediction example."),
        ("correct_1.png", "Figure 8c. Grad-CAM — additional correct example."),
        ("wrong_1.png", "Figure 8d. Grad-CAM — additional failure example."),
    ]:
        add_figure(doc, EXP / "gradcam" / name, cap, width=5.5)

    # 5 Discussion
    add_heading_custom(doc, "5. Discussion", 1)
    add_para(
        doc,
        "Three takeaways follow. First, an honest patient-level ResNet-18 on Figshare lands near "
        "macro-F1 0.95, not 0.99: still strong, but less spectacular than many slice-level reports. "
        "Second, slice-level random splits create extreme patient overlap (~99% of validation slices) "
        "and inflate metrics; journals should require patient-level protocols and overlap audits. "
        "Third, tumor masking shows a nuanced shortcut story: pretrained patient-level models lose "
        "performance when tumors are removed (they do look at lesions), yet models can re-learn high "
        "accuracy from background alone—so Grad-CAM “tumor focus” does not prove shortcut-free "
        "decision making.",
    )
    add_para(
        doc,
        "Limitations: this draft uses a single backbone (ResNet-18) and primary seed; ResNet-50 and "
        "multi-seed sensitivity remain optional. Imaging-plane metadata is incomplete on Figshare, "
        "so plane-only baselines were not fully quantified. The literature table is illustrative, "
        "not a PRISMA-style review. Figshare is 2D and saturated; Papers B–D move to 3D multi-parametric "
        "MRI, radiogenomics, survival, and segmentation robustness with overlap-safe external tests.",
    )

    # 6 Conclusion
    add_heading_custom(doc, "6. Conclusion", 1)
    add_para(
        doc,
        "Patient-level evaluation, leakage audits, and tumor-masked controls jointly show that "
        "Figshare-style 2D brain-tumor classification results are easy to overstate. We recommend "
        "frozen patient folds, mandatory overlap statistics, calibration reporting, and shortcut "
        "controls before claiming clinical relevance. Code, splits, metrics, and figures are packaged "
        "under experiments/paper_a/ in the brain-tumor-radiogenomics repository.",
    )

    # Reproducibility
    add_heading_custom(doc, "7. Reproducibility", 1)
    add_para(
        doc,
        "Repository: https://github.com/itxsamad1/brain-tumor-radiogenomics. "
        "Data DOI: 10.6084/m9.figshare.1512427. "
        "Kaggle dataset: itxsamad/btlab-figshare-paper-a. "
        "Training kernel: itxsamad/btlab-paper-a-full. "
        "Artifacts: experiments/paper_a/summary.json, fold*_resnet18/, leaky_slice_split/, "
        "tumor_masked/, gradcam/, audits/. Interactive demo: streamlit app under app/paper_a_demo.py.",
        first_line=False,
    )

    # References
    add_heading_custom(doc, "References", 1)
    refs = [
        "Cheng, J. et al. Figshare brain tumor dataset. https://doi.org/10.6084/m9.figshare.1512427",
        "He, K. et al. Deep Residual Learning for Image Recognition. CVPR 2016.",
        "Selvaraju, R.R. et al. Grad-CAM: Visual Explanations from Deep Networks. ICCV 2017.",
        "Geirhos, R. et al. Shortcut Learning in Deep Neural Networks. Nature Machine Intelligence 2020.",
        "Roberts, M. et al. Common pitfalls and recommendations for using machine learning to detect "
        "and prognosticate for COVID-19 using chest radiographs and CT scans. Nature Machine Intelligence 2021 "
        "(broader medical-ML leakage lessons).",
        "brain-tumor-radiogenomics project docs/PAPER_PUBLISH_PLAN.md (internal experimental plan).",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(f"[{i}] {ref}")
        set_run_font(run, size=10)

    # Appendix note
    add_heading_custom(doc, "Appendix A. Status relative to full publish checklist", 1)
    add_para(
        doc,
        "Completed: metrics pipeline; 5-fold ResNet-18; leaky baseline; dHash audit; Grad-CAM; "
        "tumor-masked eval+train; aggregated REPORT and this manuscript draft. "
        "Optional remaining: ResNet-50 / multi-seed; formal systematic literature table; "
        "camera-ready journal formatting and CLAIM checklist for later clinical papers (B–D).",
        first_line=False,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("Wrote", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
